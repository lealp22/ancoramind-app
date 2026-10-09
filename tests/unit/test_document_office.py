"""Deterministic DOCX/XLSX extraction (specs/021 Fase 2 / BL-35) — in-memory
only, same construction style as ``test_document_vision.py`` (no fixture
files)."""

from __future__ import annotations

import io
import zipfile

import docx
import openpyxl
import pytest

from src.agent.document_office import (
    DocxExtractionError,
    XlsxExtractionError,
    extract_docx_text_markdown,
    extract_xlsx_text_markdown,
)

_LIMITS = {"max_blocks": 100, "max_zip_entries": 200, "max_uncompressed_bytes": 50_000_000}
_XLSX_LIMITS = {**_LIMITS, "max_sheets": 20, "max_cols_per_sheet": 50}


def _docx_bytes(build: object) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _extract(content: bytes, **overrides: int) -> str:
    limits = {**_LIMITS, **overrides}
    return extract_docx_text_markdown(content, **limits)


def _xlsx_bytes(build: object) -> bytes:
    workbook = openpyxl.Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _extract_xlsx(content: bytes, **overrides: int) -> str:
    limits = {**_XLSX_LIMITS, **overrides}
    return extract_xlsx_text_markdown(content, **limits)


def test_extracts_heading_paragraph_and_table_as_markdown():
    def build(document: docx.Document) -> None:
        document.add_heading("Informe", level=1)
        document.add_paragraph("Texto normal del párrafo.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Nombre"
        table.cell(0, 1).text = "Valor"
        table.cell(1, 0).text = "Velocidad"
        table.cell(1, 1).text = "120"

    markdown = _extract(_docx_bytes(build))

    assert "# Informe" in markdown
    assert "Texto normal del párrafo." in markdown
    assert "| Nombre | Valor |" in markdown
    assert "| Velocidad | 120 |" in markdown


def test_heading_level_maps_to_hash_count():
    def build(document: docx.Document) -> None:
        document.add_heading("Nivel 2", level=2)

    markdown = _extract(_docx_bytes(build))

    assert markdown.strip() == "## Nivel 2"


def test_empty_paragraphs_are_skipped():
    def build(document: docx.Document) -> None:
        document.add_paragraph("")
        document.add_paragraph("Contenido real")
        document.add_paragraph("   ")

    markdown = _extract(_docx_bytes(build))

    assert markdown.strip() == "Contenido real"


def test_truncates_output_at_max_blocks_without_raising():
    def build(document: docx.Document) -> None:
        for i in range(20):
            document.add_paragraph(f"Párrafo {i}")

    markdown = _extract(_docx_bytes(build), max_blocks=5)

    assert "Párrafo 0" in markdown
    assert "Párrafo 4" in markdown
    assert "Párrafo 10" not in markdown
    assert "Párrafo 19" not in markdown


def test_malformed_zip_is_rejected():
    with pytest.raises(DocxExtractionError, match="docx_malformed"):
        _extract(b"this is not a zip file at all")


def test_valid_zip_without_document_xml_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a real docx")
    with pytest.raises(DocxExtractionError, match="docx_malformed"):
        _extract(buffer.getvalue())


def test_too_many_zip_entries_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<root/>")
        for i in range(5):
            archive.writestr(f"junk/{i}.txt", "x")
    with pytest.raises(DocxExtractionError, match="docx_too_many_entries"):
        _extract(buffer.getvalue(), max_zip_entries=3)


def test_zip_bomb_declared_size_is_rejected_without_inflating():
    """A highly compressible entry (small on disk, large once inflated) is
    rejected purely from the central-directory ``file_size`` header — the
    guard checks that declared metadata, it never decompresses the entry."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<root/>")
        archive.writestr("word/bomb.bin", b"\x00" * 2_000_000, zipfile.ZIP_DEFLATED)
    with pytest.raises(DocxExtractionError, match="docx_zip_bomb_suspected"):
        _extract(buffer.getvalue(), max_uncompressed_bytes=1_000_000)


def test_path_traversal_entry_name_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<root/>")
        archive.writestr("../../etc/passwd", "pwned")
    with pytest.raises(DocxExtractionError, match="docx_unsafe_entry_name"):
        _extract(buffer.getvalue())


def test_macro_project_entry_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<root/>")
        archive.writestr("word/vbaProject.bin", b"\x00\x01")
    with pytest.raises(DocxExtractionError, match="docx_macro_rejected"):
        _extract(buffer.getvalue())


def test_xml_entity_declaration_is_rejected_without_a_defusedxml_dependency():
    """OOXML never legitimately declares a DTD/entity; this is the guard
    against XML entity-expansion ("billion laughs") attacks (specs/021)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<!DOCTYPE root [<!ENTITY lol "lol">]><root>&lol;</root>',
        )
    with pytest.raises(DocxExtractionError, match="docx_xml_entity_rejected"):
        _extract(buffer.getvalue())


# ── XLSX ──────────────────────────────────────────────────────────────────


def test_extracts_sheet_as_markdown_table_with_sheet_heading():
    def build(workbook: openpyxl.Workbook) -> None:
        sheet = workbook.active
        sheet.title = "Gastos"
        sheet.append(["Concepto", "Importe"])
        sheet.append(["Taller", 120])

    markdown = _extract_xlsx(_xlsx_bytes(build))

    assert "## Gastos" in markdown
    assert "| Concepto | Importe |" in markdown
    assert "| Taller | 120 |" in markdown


def test_xlsx_formulas_render_the_cached_value_not_the_formula():
    def build(workbook: openpyxl.Workbook) -> None:
        sheet = workbook.active
        sheet.append(["A", "B", "Total"])
        sheet["A2"] = 2
        sheet["B2"] = 3
        sheet["C2"] = "=A2+B2"

    markdown = _extract_xlsx(_xlsx_bytes(build))

    # A freshly-written workbook has no cached formula result (Excel/LibreOffice
    # never computed it), so openpyxl's data_only read returns None for it —
    # the key guarantee is that the literal formula string never appears.
    assert "=A2+B2" not in markdown


def test_xlsx_truncates_rows_across_sheets_by_shared_block_budget():
    def build(workbook: openpyxl.Workbook) -> None:
        sheet = workbook.active
        for i in range(20):
            sheet.append([f"fila {i}"])

    markdown = _extract_xlsx(_xlsx_bytes(build), max_blocks=5)

    assert "fila 0" in markdown
    assert "fila 4" in markdown
    assert "fila 10" not in markdown
    assert "fila 19" not in markdown


def test_xlsx_blank_rows_still_consume_the_shared_budget():
    """An adversarial sheet of blank rows must not iterate unbounded."""

    def build(workbook: openpyxl.Workbook) -> None:
        sheet = workbook.active
        for _ in range(50):
            sheet.append([None, None])
        sheet.append(["nunca llega", "porque el budget ya se agotó"])

    markdown = _extract_xlsx(_xlsx_bytes(build), max_blocks=10)

    assert markdown.strip() == ""


def test_xlsx_respects_max_sheets():
    def build(workbook: openpyxl.Workbook) -> None:
        workbook.active.title = "Uno"
        workbook.active.append(["x"])
        workbook.create_sheet("Dos").append(["y"])
        workbook.create_sheet("Tres").append(["z"])

    markdown = _extract_xlsx(_xlsx_bytes(build), max_sheets=1)

    assert "## Uno" in markdown
    assert "## Dos" not in markdown
    assert "## Tres" not in markdown


def test_xlsx_respects_max_cols_per_sheet():
    def build(workbook: openpyxl.Workbook) -> None:
        workbook.active.append(["a", "b", "c", "d"])

    markdown = _extract_xlsx(_xlsx_bytes(build), max_cols_per_sheet=2)

    assert "| a | b |" in markdown
    assert "c" not in markdown
    assert "d" not in markdown


def test_xlsx_malformed_zip_is_rejected():
    with pytest.raises(XlsxExtractionError, match="xlsx_malformed"):
        _extract_xlsx(b"this is not a zip file at all")


def test_xlsx_zip_bomb_declared_size_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<root/>")
        archive.writestr("xl/bomb.bin", b"\x00" * 2_000_000, zipfile.ZIP_DEFLATED)
    with pytest.raises(XlsxExtractionError, match="xlsx_zip_bomb_suspected"):
        _extract_xlsx(buffer.getvalue(), max_uncompressed_bytes=1_000_000)


def test_xlsx_path_traversal_entry_name_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<root/>")
        archive.writestr("../../etc/passwd", "pwned")
    with pytest.raises(XlsxExtractionError, match="xlsx_unsafe_entry_name"):
        _extract_xlsx(buffer.getvalue())


def test_xlsx_macro_project_entry_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<root/>")
        archive.writestr("xl/vbaProject.bin", b"\x00\x01")
    with pytest.raises(XlsxExtractionError, match="xlsx_macro_rejected"):
        _extract_xlsx(buffer.getvalue())
